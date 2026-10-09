import requests
import urllib3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from getpass import getpass

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ==============================
# F5 CONNECTION DETAILS
# ==============================
F5_HOST = "10.225.53.6"
USERNAME = "admin"
PASSWORD = getpass("Password: ")

POLICY_NAME = "HTTPS_Redirect_Policy_v1"
OUTPUT_FILE = "F5_Policy_Report.xlsx"

BASE_URL = f"https://{F5_HOST}/mgmt/tm"


# ==============================
# REST REQUEST FUNCTION
# ==============================
def f5_get(uri):
    response = requests.get(
        BASE_URL + uri,
        auth=(USERNAME, PASSWORD),
        verify=False,
        timeout=30
    )

    response.raise_for_status()
    return response.json()


# ==============================
# CONVERT F5 REST PATH
# ==============================
def convert_selflink_to_uri(selflink):
    """
    Convert F5 REST selfLink into a URI that can
    be used with BASE_URL.
    """

    if "/mgmt/tm" in selflink:
        return selflink.split("/mgmt/tm", 1)[1].split("?")[0]

    return selflink


# ==============================
# GET POLICY
# ==============================
def get_policy(policy_name):

    print(f"\nGetting policy: {policy_name}")

    uri = (
        f"/ltm/policy/~Common~{policy_name}"
        "?expandSubcollections=true"
    )

    return f5_get(uri)


# ==============================
# FORMAT CONDITION
# ==============================
def format_condition(condition):

    details = []

    for key in [
        "httpHost",
        "httpUri",
        "httpMethod",
        "tcp",
        "sslExtension",
        "sslClientHello",
        "serverName",
        "host",
        "path",
        "pathSegment",
        "scheme",
        "address",
        "port",
        "external"
    ]:
        if condition.get(key):
            details.append(key)

    if condition.get("operator"):
        details.append(f"operator={condition['operator']}")

    if condition.get("values"):
        details.append(
            "values=" + ", ".join(
                str(v) for v in condition["values"]
            )
        )

    if condition.get("all"):
        details.append("all=true")

    if condition.get("caseInsensitive"):
        details.append("case-insensitive=true")

    return "; ".join(details)


# ==============================
# DETECT ACTION TYPE/TARGET
# ==============================
def parse_action(action):

    action_type = ""
    forward_type = ""
    target = ""
    redirect = ""

    # Pool forwarding
    if action.get("pool"):
        action_type = "Forward"
        forward_type = "Pool"
        target = action.get("pool")

    # Virtual Server forwarding
    elif action.get("virtual"):
        action_type = "Forward"
        forward_type = "Virtual Server"
        target = action.get("virtual")

    # Node forwarding
    elif action.get("node"):
        action_type = "Forward"
        forward_type = "Node"
        target = action.get("node")

    # Redirect
    elif action.get("location"):
        action_type = "Redirect"
        forward_type = "HTTP Redirect"
        redirect = action.get("location")

    # Reset
    elif action.get("reset"):
        action_type = "Reset"

    # Drop
    elif action.get("drop"):
        action_type = "Drop"

    # Disable
    elif action.get("disable"):
        action_type = "Disable"

    else:
        # Preserve useful information for action types
        # not explicitly handled above.
        action_type = ", ".join(
            key
            for key, value in action.items()
            if value is True
        )

    return action_type, forward_type, target, redirect


# ==============================
# GET COLLECTION ITEMS
# ==============================
def get_collection(parent, field_name):

    collection = parent.get(field_name, {})

    # Already expanded
    items = collection.get("items", [])

    if items:
        return items

    # Retrieve through REST link
    link = collection.get("link")

    if link:
        uri = convert_selflink_to_uri(link)
        data = f5_get(uri)
        return data.get("items", [])

    return []


# ==============================
# MAIN
# ==============================
def main():

    policy = get_policy(POLICY_NAME)

    print(f"Policy found: {policy.get('name')}")

    rules = get_collection(policy, "rulesReference")

    print(f"Rules found: {len(rules)}")

    report_rows = []

    for rule in rules:

        rule_name = rule.get("name", "")
        ordinal = rule.get("ordinal", "")

        print(f"Processing rule: {rule_name}")

        conditions = get_collection(
            rule,
            "conditionsReference"
        )

        actions = get_collection(
            rule,
            "actionsReference"
        )

        # ------------------------------------------------
        # Build condition description
        # ------------------------------------------------

        condition_text = []
        operands = []
        events = []
        values = []

        for condition in conditions:

            formatted = format_condition(condition)

            if formatted:
                condition_text.append(formatted)

            operand = condition.get("operand", "")

            if operand:
                operands.append(str(operand))

            event = condition.get("event", "")

            if event:
                events.append(str(event))

            condition_values = condition.get("values", [])

            if condition_values:
                values.extend(
                    str(v) for v in condition_values
                )

        condition_string = "\n".join(condition_text)
        operand_string = "\n".join(operands)
        event_string = "\n".join(events)
        value_string = "\n".join(values)

        # Rule may theoretically have no actions
        if not actions:

            report_rows.append([
                POLICY_NAME,
                rule_name,
                ordinal,
                condition_string,
                operand_string,
                event_string,
                value_string,
                "",
                "",
                "",
                ""
            ])

            continue

        # ------------------------------------------------
        # Process actions
        # ------------------------------------------------

        for action in actions:

            (
                action_type,
                forward_type,
                target,
                redirect
            ) = parse_action(action)

            report_rows.append([
                POLICY_NAME,
                rule_name,
                ordinal,
                condition_string,
                operand_string,
                event_string,
                value_string,
                action_type,
                forward_type,
                target,
                redirect
            ])

    # ==============================
    # CREATE EXCEL
    # ==============================

    wb = Workbook()

    ws = wb.active
    ws.title = "F5 Policy"

    headers = [
        "Policy Name",
        "Rule Name",
        "Rule Ordinal",
        "Condition",
        "Operand",
        "Event",
        "Values",
        "Action",
        "Forward Type",
        "Target",
        "Redirect Location"
    ]

    ws.append(headers)

    # Header formatting
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Data
    for row in report_rows:
        ws.append(row)

    # Formatting
    widths = {
        "A": 30,
        "B": 35,
        "C": 15,
        "D": 65,
        "E": 25,
        "F": 25,
        "G": 45,
        "H": 20,
        "I": 25,
        "J": 45,
        "K": 60
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_FILE)

    print("\n================================")
    print("Report generated successfully")
    print("================================")
    print(f"File: {OUTPUT_FILE}")
    print(f"Policy: {POLICY_NAME}")
    print(f"Rules: {len(rules)}")
    print(f"Report rows: {len(report_rows)}")


if __name__ == "__main__":
    main()
