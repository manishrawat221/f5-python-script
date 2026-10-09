#!/usr/bin/env python3

import requests
import pandas as pd
import urllib3
from requests.auth import HTTPBasicAuth
from getpass import getpass
from openpyxl.styles import Alignment, Font, PatternFill

# Disable SSL certificate warnings
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


# ============================================================
# F5 Connection Details
# ============================================================

F5_HOST = "10.79.53.6"
USERNAME = "admin"
PASSWORD = getpass("Enter F5 Password: ")

auth = HTTPBasicAuth(USERNAME, PASSWORD)


# ============================================================
# Helper Function - API GET
# ============================================================

def api_get(url):

    response = requests.get(
        url,
        auth=auth,
        verify=False,
        timeout=30
    )

    response.raise_for_status()

    return response.json()


# ============================================================
# Helper Function - Convert F5 REST Reference URL
# ============================================================

def get_reference_url(link):

    if not link:
        return ""

    return link.replace(
        "https://localhost",
        f"https://{F5_HOST}"
    )


# ============================================================
# Helper Function - Build BIG-IP REST Object Path
# ============================================================

def rest_object_path(partition, name):

    partition = partition or "Common"

    return f"~{partition}~{name}"


# ============================================================
# Get Virtual Servers
# ============================================================

virtuals_url = (
    f"https://{F5_HOST}/mgmt/tm/ltm/virtual"
    "?$select=name,partition,destination,pool,rules,"
    "description,policiesReference,trafficMatchingCriteria"
)

try:

    virtuals = api_get(
        virtuals_url
    ).get("items", [])

except Exception as err:

    print()
    print("ERROR: Unable to retrieve Virtual Servers")
    print(err)

    raise SystemExit(1)


print()
print("==============================================")
print("F5 VIP Discovery")
print("==============================================")
print(f"F5 Device       : {F5_HOST}")
print(f"Virtual Servers : {len(virtuals)}")
print("==============================================")
print()


report_rows = []


# ============================================================
# Process VIPs
# ============================================================

for vip in virtuals:

    vip_name = vip.get(
        "name",
        ""
    )

    vip_partition = vip.get(
        "partition",
        "Common"
    )

    vip_ip = vip.get(
        "destination",
        ""
    )

    vip_description = vip.get(
        "description",
        ""
    )

    print(f"Processing VIP: {vip_name}")


    # ========================================================
    # Pool Information
    # ========================================================

    pool_full = vip.get(
        "pool",
        ""
    )

    pool_name = ""
    pool_partition = "Common"

    if pool_full:

        pool_parts = pool_full.strip("/").split("/")

        if len(pool_parts) >= 2:

            pool_partition = pool_parts[-2]
            pool_name = pool_parts[-1]

        else:

            pool_name = pool_parts[-1]


    # ========================================================
    # iRules
    # ========================================================

    irule_names = []

    for rule in vip.get("rules", []):

        rule_name = rule.split("/")[-1]

        if rule_name:
            irule_names.append(rule_name)

    irules = ", ".join(irule_names)


    # ========================================================
    # Policies
    # ========================================================

    policy_names = []

    try:

        policies_ref = vip.get(
            "policiesReference",
            {}
        )

        # ----------------------------------------------------
        # Case 1: Policy objects already returned
        # ----------------------------------------------------

        policy_items = policies_ref.get(
            "items",
            []
        )

        for policy in policy_items:

            policy_name = policy.get(
                "name",
                ""
            )

            if policy_name:
                policy_names.append(
                    policy_name
                )


        # ----------------------------------------------------
        # Case 2: Only policiesReference link returned
        # ----------------------------------------------------

        if (
            not policy_names
            and policies_ref.get("link")
        ):

            policy_url = get_reference_url(
                policies_ref.get("link")
            )

            policy_data = api_get(
                policy_url
            )

            for policy in policy_data.get(
                "items",
                []
            ):

                policy_name = policy.get(
                    "name",
                    ""
                )

                if policy_name:

                    policy_names.append(
                        policy_name
                    )

    except Exception as err:

        policy_names.append(
            f"ERROR: {str(err)}"
        )


    # Remove duplicate policy names

    policy_names = list(
        dict.fromkeys(policy_names)
    )

    policies = ", ".join(
        policy_names
    )


    # ========================================================
    # Traffic Matching Criteria
    # ========================================================

    traffic_matching_criteria = vip.get(
        "trafficMatchingCriteria",
        ""
    )

    if traffic_matching_criteria:

        traffic_matching_criteria = (
            traffic_matching_criteria
            .split("/")[-1]
        )


    # ========================================================
    # VIP Status
    # ========================================================

    vip_status = "Unknown"

    try:

        vip_rest_path = rest_object_path(
            vip_partition,
            vip_name
        )

        stats_url = (
            f"https://{F5_HOST}"
            f"/mgmt/tm/ltm/virtual/"
            f"{vip_rest_path}/stats"
        )

        stats = api_get(
            stats_url
        )

        entries = stats.get(
            "entries",
            {}
        )

        for entry in entries.values():

            nested = (
                entry
                .get("nestedStats", {})
                .get("entries", {})
            )

            availability = nested.get(
                "status.availabilityState",
                {}
            )

            if availability:

                vip_status = availability.get(
                    "description",
                    "Unknown"
                )

                break

    except Exception as err:

        vip_status = f"ERROR: {str(err)}"


    # ========================================================
    # Pool Members
    # ========================================================

    pool_member_ips = []
    member_states = []


    if pool_name:

        try:

            pool_rest_path = rest_object_path(
                pool_partition,
                pool_name
            )

            members_url = (
                f"https://{F5_HOST}"
                f"/mgmt/tm/ltm/pool/"
                f"{pool_rest_path}/members"
            )

            pool_members = (
                api_get(
                    members_url
                )
                .get("items", [])
            )


            for member in pool_members:

                member_ip = member.get(
                    "address",
                    ""
                )

                member_state = member.get(
                    "state",
                    ""
                )

                if member_ip:

                    pool_member_ips.append(
                        member_ip
                    )

                    member_states.append(
                        f"{member_ip} ({member_state})"
                    )


        except Exception as err:

            pool_member_ips.append(
                f"ERROR: {str(err)}"
            )


    # ========================================================
    # Console Output
    # ========================================================

    print(f"  Status : {vip_status}")
    print(f"  Pool   : {pool_name}")
    print(f"  iRules : {irules}")
    print(f"  Policy : {policies}")
    print(
        f"  TMC    : "
        f"{traffic_matching_criteria}"
    )

    print()


    # ========================================================
    # Add One Row Per VIP
    # ========================================================

    report_rows.append(
        {
            "VIP_Name":
                vip_name,

            "VIP_Description":
                vip_description,

            "VIP_IP":
                vip_ip,

            "VIP_Status":
                vip_status,

            "Pool_Name":
                pool_name,

            "Pool_Member_Count":
                len(pool_member_ips),

            "Pool_Members":
                ", ".join(
                    pool_member_ips
                ),

            "Member_States":
                ", ".join(
                    member_states
                ),

            "iRules":
                irules,

            "Policy":
                policies,

            "Traffic_Matching_Criteria":
                traffic_matching_criteria
        }
    )


# ============================================================
# Create DataFrame
# ============================================================

columns = [
    "VIP_Name",
    "VIP_Description",
    "VIP_IP",
    "VIP_Status",
    "Pool_Name",
    "Pool_Member_Count",
    "Pool_Members",
    "Member_States",
    "iRules",
    "Policy",
    "Traffic_Matching_Criteria"
]

df = pd.DataFrame(
    report_rows,
    columns=columns
)


# ============================================================
# Export To Excel
# ============================================================

excel_file = "F5_VIP_Report.xlsx"

with pd.ExcelWriter(
    excel_file,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="VIP_Report",
        index=False
    )

    worksheet = writer.sheets["VIP_Report"]

    # ========================================================
    # Column Widths
    # ========================================================

    worksheet.column_dimensions["A"].width = 35
    worksheet.column_dimensions["B"].width = 50
    worksheet.column_dimensions["C"].width = 30
    worksheet.column_dimensions["D"].width = 20
    worksheet.column_dimensions["E"].width = 35
    worksheet.column_dimensions["F"].width = 20
    worksheet.column_dimensions["G"].width = 60
    worksheet.column_dimensions["H"].width = 70
    worksheet.column_dimensions["I"].width = 50
    worksheet.column_dimensions["J"].width = 45
    worksheet.column_dimensions["K"].width = 55

    # ========================================================
    # Header Formatting
    # ========================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    # Format header row
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment (
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
    #=====================================================
    # Data Row Formatting
    # ========================================================

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Freeze header
    worksheet.freeze_panes = "A2"

    # Enable filters
    worksheet.auto_filter.ref = worksheet.dimensions



# ============================================================
# Completed
# ============================================================

print()
print("==============================================")
print("F5 VIP Report Generated Successfully")
print("==============================================")
print(f"F5 Device       : {F5_HOST}")
print(f"Virtual Servers : {len(df)}")
print(f"Excel File      : {excel_file}")
print("==============================================")
